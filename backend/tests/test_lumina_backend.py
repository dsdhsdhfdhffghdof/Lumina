"""Comprehensive backend tests for Lumina Academic API.

Covers auth (synthetic sessions), materials CRUD, upload validation,
cross-user isolation, retry, tutor grounded chat + fallback, dashboard,
onboarding, logout, expired session, and CORS preflight.
"""
import io
import os
import uuid
import asyncio
from datetime import datetime, timezone, timedelta
from pathlib import Path

import pytest
import httpx
from motor.motor_asyncio import AsyncIOMotorClient
from dotenv import load_dotenv
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document as DocxDocument

load_dotenv("/app/backend/.env")

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://source-tutor.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]

TEST_EMAIL_SUFFIX = "@lumina-test.local"

# ---------- helpers ----------

def build_test_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setFont("Helvetica", 14)
    c.drawString(72, 720, "Lumina Test Textbook - Cellular Respiration")
    c.drawString(72, 690, "Chapter 1 introduces cellular respiration as the biochemical process")
    c.drawString(72, 672, "in which cells convert glucose into usable energy stored as ATP.")
    c.showPage()
    c.setFont("Helvetica", 14)
    c.drawString(72, 720, "Chapter 2 - The Electron Transport Chain")
    c.drawString(72, 690, "The electron transport chain pumps protons across the inner")
    c.drawString(72, 672, "mitochondrial membrane, creating an electrochemical gradient.")
    c.drawString(72, 654, "ATP synthase uses this gradient to produce ATP from ADP and Pi.")
    c.showPage()
    c.save()
    return buf.getvalue()


def build_test_docx() -> bytes:
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Photosynthesis is the process by which green plants convert sunlight into chemical energy.")
    doc.add_paragraph("Chlorophyll absorbs light in the blue and red regions of the spectrum.")
    for _ in range(30):
        doc.add_paragraph("This is filler content to force chunking into multiple logical pages for Lumina testing.")
    doc.save(buf)
    return buf.getvalue()


def build_test_txt(label: str = "txt") -> bytes:
    lines = [f"Introduction to Quantum Entanglement ({label})"]
    for i in range(80):
        lines.append(f"Line {i}: Entangled particles share a quantum state regardless of distance.")
    return ("\n".join(lines)).encode("utf-8")


# ---------- fixtures ----------

@pytest.fixture(scope="session")
def event_loop():
    loop = asyncio.new_event_loop()
    yield loop
    loop.close()


async def _mint_session(email: str, hours: int = 4):
    cli = AsyncIOMotorClient(MONGO_URL)
    db = cli[DB_NAME]
    uid = "user_" + uuid.uuid4().hex[:12]
    tok = uuid.uuid4().hex
    now = datetime.now(timezone.utc).isoformat()
    exp = (datetime.now(timezone.utc) + timedelta(hours=hours)).isoformat()
    await db.users.insert_one({
        "user_id": uid, "google_id": "test_" + uid, "email": email,
        "name": "Test " + email.split("@")[0], "picture": "",
        "created_at": now, "first_sign_in_at": now, "last_sign_in_at": now,
        "onboarded": False, "preferred_language": "English",
    })
    await db.user_sessions.insert_one({
        "user_id": uid, "session_token": tok, "created_at": now, "expires_at": exp,
    })
    cli.close()
    return uid, tok


async def _mint_expired_session(email: str):
    cli = AsyncIOMotorClient(MONGO_URL)
    db = cli[DB_NAME]
    uid = "user_" + uuid.uuid4().hex[:12]
    tok = uuid.uuid4().hex
    now = datetime.now(timezone.utc).isoformat()
    exp = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
    await db.users.insert_one({
        "user_id": uid, "email": email, "name": "Expired", "created_at": now,
        "first_sign_in_at": now, "last_sign_in_at": now, "onboarded": False,
    })
    await db.user_sessions.insert_one({
        "user_id": uid, "session_token": tok, "created_at": now, "expires_at": exp,
    })
    cli.close()
    return uid, tok


@pytest.fixture(scope="module")
def user_a():
    email = f"a-{uuid.uuid4().hex[:6]}{TEST_EMAIL_SUFFIX}"
    uid, tok = asyncio.get_event_loop().run_until_complete(_mint_session(email))
    return {"user_id": uid, "token": tok, "email": email}


@pytest.fixture(scope="module")
def user_b():
    email = f"b-{uuid.uuid4().hex[:6]}{TEST_EMAIL_SUFFIX}"
    uid, tok = asyncio.get_event_loop().run_until_complete(_mint_session(email))
    return {"user_id": uid, "token": tok, "email": email}


@pytest.fixture(scope="module")
def http():
    with httpx.Client(base_url=API, timeout=90) as c:
        yield c


@pytest.fixture(scope="module", autouse=True)
def cleanup_at_end(user_a, user_b):
    yield
    async def _cleanup():
        cli = AsyncIOMotorClient(MONGO_URL)
        db = cli[DB_NAME]
        ids = [user_a["user_id"], user_b["user_id"]]
        await db.users.delete_many({"user_id": {"$in": ids}})
        await db.user_sessions.delete_many({"user_id": {"$in": ids}})
        await db.materials.delete_many({"user_id": {"$in": ids}})
        await db.material_pages.delete_many({"user_id": {"$in": ids}})
        await db.chat_messages.delete_many({"user_id": {"$in": ids}})
        # expired/logout test users (prefixed with our suffix)
        test_users = await db.users.find({"email": {"$regex": TEST_EMAIL_SUFFIX}}, {"user_id": 1}).to_list(length=500)
        test_uids = [u["user_id"] for u in test_users]
        if test_uids:
            await db.user_sessions.delete_many({"user_id": {"$in": test_uids}})
            await db.users.delete_many({"user_id": {"$in": test_uids}})
            await db.materials.delete_many({"user_id": {"$in": test_uids}})
            await db.material_pages.delete_many({"user_id": {"$in": test_uids}})
            await db.chat_messages.delete_many({"user_id": {"$in": test_uids}})
        cli.close()
    asyncio.get_event_loop().run_until_complete(_cleanup())


def auth(user):
    return {"Authorization": f"Bearer {user['token']}"}


# ---------- auth tests ----------

class TestAuth:
    def test_unauth_materials_401(self, http):
        r = http.get("/materials")
        assert r.status_code == 401

    def test_unauth_me_401(self, http):
        r = http.get("/auth/me")
        assert r.status_code == 401

    def test_invalid_session_id_401(self, http):
        r = http.post("/auth/session", json={"session_id": "definitely-not-valid-xyz"})
        assert r.status_code == 401

    def test_synthetic_session_me_200(self, http, user_a):
        r = http.get("/auth/me", headers=auth(user_a))
        assert r.status_code == 200
        data = r.json()
        assert data["email"] == user_a["email"]
        assert data["user_id"] == user_a["user_id"]

    def test_expired_session_rejected(self, http):
        email = f"exp-{uuid.uuid4().hex[:6]}{TEST_EMAIL_SUFFIX}"
        uid, tok = asyncio.get_event_loop().run_until_complete(_mint_expired_session(email))
        r = http.get("/auth/me", headers={"Authorization": f"Bearer {tok}"})
        assert r.status_code == 401


# ---------- upload / materials ----------

class TestUploadValidation:
    def test_unsupported_extension_400(self, http, user_a):
        r = http.post("/materials/upload",
                      files={"file": ("x.exe", b"MZ...", "application/octet-stream")},
                      headers=auth(user_a))
        assert r.status_code == 400

    def test_unsupported_zip_400(self, http, user_a):
        r = http.post("/materials/upload",
                      files={"file": ("a.zip", b"PK\x03\x04", "application/zip")},
                      headers=auth(user_a))
        assert r.status_code == 400

    def test_unsupported_gif_400(self, http, user_a):
        # PNG/JPG/WEBP/TIFF/BMP are now accepted (OCR path); GIF is still rejected.
        r = http.post("/materials/upload",
                      files={"file": ("a.gif", b"GIF89a\x00\x00", "image/gif")},
                      headers=auth(user_a))
        assert r.status_code == 400

    def test_empty_file_400(self, http, user_a):
        r = http.post("/materials/upload",
                      files={"file": ("empty.pdf", b"", "application/pdf")},
                      headers=auth(user_a))
        assert r.status_code == 400

    def test_oversize_file_413(self, http, user_a):
        big = b"a" * (25 * 1024 * 1024 + 100)
        r = http.post("/materials/upload",
                      files={"file": ("big.txt", big, "text/plain")},
                      headers=auth(user_a))
        assert r.status_code == 413


# ---------- end-to-end flow ----------

class TestMaterialsAndTutor:
    pdf_id = None
    docx_id = None
    txt_id = None

    def test_upload_pdf(self, http, user_a):
        pdf = build_test_pdf()
        r = http.post("/materials/upload",
                      files={"file": ("textbook.pdf", pdf, "application/pdf")},
                      headers=auth(user_a))
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["status"] == "ready", f"status={data['status']} err={data.get('error')}"
        assert data["page_count"] >= 1
        assert data["type"] == "PDF"
        TestMaterialsAndTutor.pdf_id = data["id"]

        # verify DB persistence
        async def _check():
            cli = AsyncIOMotorClient(MONGO_URL)
            db = cli[DB_NAME]
            m = await db.materials.find_one({"material_id": data["id"]})
            pages = await db.material_pages.count_documents({"material_id": data["id"]})
            cli.close()
            return m, pages
        m, pages = asyncio.get_event_loop().run_until_complete(_check())
        assert m is not None
        assert pages >= 1

    def test_upload_docx(self, http, user_a):
        docx = build_test_docx()
        r = http.post("/materials/upload",
                      files={"file": ("notes.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                      headers=auth(user_a))
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["status"] == "ready", data.get("error")
        assert data["page_count"] >= 1
        TestMaterialsAndTutor.docx_id = data["id"]

    def test_upload_txt(self, http, user_a):
        txt = build_test_txt()
        r = http.post("/materials/upload",
                      files={"file": ("notes.txt", txt, "text/plain")},
                      headers=auth(user_a))
        assert r.status_code == 200
        data = r.json()
        assert data["status"] == "ready"
        assert data["page_count"] >= 1
        TestMaterialsAndTutor.txt_id = data["id"]

    def test_upload_md(self, http, user_a):
        md = b"# Title\n\n" + build_test_txt("md")
        r = http.post("/materials/upload",
                      files={"file": ("notes.md", md, "text/markdown")},
                      headers=auth(user_a))
        assert r.status_code == 200
        data = r.json()
        assert data["status"] == "ready"
        assert data["page_count"] >= 1

    def test_list_only_owner(self, http, user_a, user_b):
        r = http.get("/materials", headers=auth(user_a))
        assert r.status_code == 200
        a_items = r.json()["items"]
        assert len(a_items) >= 4
        assert all("filename" in i for i in a_items)

        r = http.get("/materials", headers=auth(user_b))
        assert r.status_code == 200
        assert r.json()["count"] == 0

    def test_download_isolation(self, http, user_a, user_b):
        mid = TestMaterialsAndTutor.pdf_id
        # unauth 401
        assert http.get(f"/materials/{mid}/download").status_code == 401
        # other user 404
        r = http.get(f"/materials/{mid}/download", headers=auth(user_b))
        assert r.status_code == 404
        # owner 200
        r = http.get(f"/materials/{mid}/download", headers=auth(user_a))
        assert r.status_code == 200
        assert len(r.content) > 100

    def test_onboarding_true_after_ready(self, http, user_a):
        r = http.get("/onboarding/status", headers=auth(user_a))
        assert r.status_code == 200
        data = r.json()
        assert data["onboarded"] is True
        assert data["ready_materials"] >= 1

    def test_onboarding_false_for_user_b(self, http, user_b):
        r = http.get("/onboarding/status", headers=auth(user_b))
        assert r.status_code == 200
        assert r.json()["onboarded"] is False

    def test_tutor_grounded_answer(self, http, user_a):
        r = http.post("/tutor/chat",
                      json={"message": "What does ATP synthase use to produce ATP?"},
                      headers=auth(user_a), timeout=90)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["grounded"] is True, f"answer={data['answer'][:200]}"
        assert len(data["citations"]) > 0
        # Must contain at least one [filename, p.N] citation
        import re
        assert re.search(r"\[[^\[\]]+,\s*p\.\d+\]", data["answer"]), f"no citation marker in: {data['answer'][:300]}"

    def test_tutor_off_material_fallback(self, http, user_a):
        r = http.post("/tutor/chat",
                      json={"message": "What is the capital of France?"},
                      headers=auth(user_a), timeout=90)
        assert r.status_code == 200
        data = r.json()
        assert data["grounded"] is False
        assert data["answer"] == "I couldn't find this information in your uploaded materials. Please upload the relevant page or document."

    def test_tutor_no_materials_fallback(self, http, user_b):
        r = http.post("/tutor/chat",
                      json={"message": "What is ATP?"},
                      headers=auth(user_b), timeout=90)
        assert r.status_code == 200
        data = r.json()
        assert data["grounded"] is False
        assert data["sources_available"] == 0
        assert data["answer"] == "I couldn't find this information in your uploaded materials. Please upload the relevant page or document."

    def test_history_user_scoped(self, http, user_a, user_b):
        ra = http.get("/tutor/history", headers=auth(user_a)).json()
        rb = http.get("/tutor/history", headers=auth(user_b)).json()
        assert len(ra["items"]) >= 2
        assert all(it["user_id"] == user_a["user_id"] for it in ra["items"])
        # user_b should only have its own records (the no-materials one above)
        assert all(it["user_id"] == user_b["user_id"] for it in rb["items"])
        # no cross-leak
        a_ids = {it["id"] for it in ra["items"]}
        b_ids = {it["id"] for it in rb["items"]}
        assert a_ids.isdisjoint(b_ids)

    def test_dashboard_scoped(self, http, user_a):
        r = http.get("/dashboard", headers=auth(user_a))
        assert r.status_code == 200
        data = r.json()
        assert data["ready_count"] >= 4
        assert data["total_pages"] >= 4
        assert data["questions"] >= 2
        assert isinstance(data["grounded"], int)
        assert 0 <= data["confidence"] <= 100

    def test_retry_reprocesses(self, http, user_a):
        """Simulate a failed material by flipping status; retry should re-ready it."""
        mid = TestMaterialsAndTutor.docx_id
        async def _fail():
            cli = AsyncIOMotorClient(MONGO_URL)
            db = cli[DB_NAME]
            await db.materials.update_one({"material_id": mid}, {"$set": {"status": "failed", "error": "simulated"}})
            cli.close()
        asyncio.get_event_loop().run_until_complete(_fail())
        r = http.post(f"/materials/{mid}/retry", headers=auth(user_a))
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["status"] == "ready"

    def test_delete_by_owner(self, http, user_a, user_b):
        mid = TestMaterialsAndTutor.txt_id
        # other user delete returns 404
        assert http.delete(f"/materials/{mid}", headers=auth(user_b)).status_code == 404
        # owner delete 200
        r = http.delete(f"/materials/{mid}", headers=auth(user_a))
        assert r.status_code == 200
        # list excludes it
        items = http.get("/materials", headers=auth(user_a)).json()["items"]
        assert all(i["id"] != mid for i in items)


# ---------- logout + CORS ----------

class TestLogoutAndCors:
    def test_logout_invalidates_session(self, http):
        # mint a one-off session, logout, verify 401
        email = f"lo-{uuid.uuid4().hex[:6]}{TEST_EMAIL_SUFFIX}"
        uid, tok = asyncio.get_event_loop().run_until_complete(_mint_session(email))
        h = {"Authorization": f"Bearer {tok}"}
        assert http.get("/auth/me", headers=h).status_code == 200
        assert http.post("/auth/logout", headers=h).status_code == 200
        assert http.get("/auth/me", headers=h).status_code == 401

    def test_cors_preflight(self, http):
        r = http.request("OPTIONS", "/materials", headers={
            "Origin": "https://example.com",
            "Access-Control-Request-Method": "GET",
            "Access-Control-Request-Headers": "authorization",
        })
        assert r.status_code in (200, 204)
        assert r.headers.get("access-control-allow-credentials", "").lower() == "true"
        # origin reflected (allow_origin_regex=.*)
        assert r.headers.get("access-control-allow-origin") in ("https://example.com", "*")
